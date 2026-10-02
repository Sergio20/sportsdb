#!/usr/bin/env python3
"""Exporta la base de datos SQLite a un libro Excel (una hoja por tabla/vista).

    python export_excel.py [--db data/deportes.db] [--out data/deportes.xlsx]
"""
from __future__ import annotations

import argparse
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from sportsdb.common import DEFAULT_DB

FONT = Font(name="Arial", size=10)
HEAD_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", start_color="1F3864")

SHEETS = [
    ("Partidos", "Todos los partidos de las tres ligas (uno por fila).",
     "SELECT match_id, league, season, date, time, time_ref, phase, round, home_team, away_team, "
     "home_score, away_score, result, status, overtimes, venue, attendance, referees, "
     "home_coach, away_coach FROM matches ORDER BY league, date, match_id"),
    ("Parciales", "Puntos/goles por periodo: fútbol 1T y 2T; baloncesto Q1-Q4 y prórrogas (OT).",
     "SELECT p.match_id, m.league, m.season, m.date, m.home_team, m.away_team, p.period, "
     "p.label, p.home, p.away FROM periods p JOIN matches m USING(match_id) "
     "ORDER BY m.league, m.date, p.match_id, p.period"),
    ("LaLiga", "LaLiga: partido + descanso + estadísticas + cuotas principales (pre y cierre).",
     "SELECT * FROM v_laliga ORDER BY date, match_id"),
    ("Baloncesto", "Euroliga y Liga Endesa: partido + parciales + estadísticas de equipo.",
     "SELECT * FROM v_baloncesto ORDER BY league, date, match_id"),
    ("Basket_jugadores", "Estadísticas individuales por partido (Euroliga y Liga Endesa).",
     "SELECT s.match_id, m.league, m.season, m.date, s.team, s.is_home, s.player_id, s.player, "
     "s.starter, ROUND(s.seconds / 60.0, 1) AS minutes, s.points, s.fg2m, s.fg2a, s.fg3m, s.fg3a, "
     "s.ftm, s.fta, s.oreb, s.dreb, s.treb, s.assists, s.steals, s.turnovers, s.blocks, "
     "s.blocks_against, s.fouls, s.fouls_drawn, s.dunks, s.plus_minus, s.rating "
     "FROM basket_player_stats s JOIN matches m USING(match_id) "
     "ORDER BY m.league, m.date, s.match_id, s.is_home DESC, s.seconds DESC"),
    ("Cuotas_LaLiga", "Todas las cuotas de LaLiga en formato largo (casa, mercado, selección, etapa).",
     "SELECT o.match_id, m.season, m.date, m.home_team, m.away_team, o.bookmaker, b.name AS "
     "bookmaker_name, o.market, o.selection, o.stage, o.line, o.price FROM odds o "
     "JOIN matches m USING(match_id) LEFT JOIN bookmakers b ON b.code = o.bookmaker "
     "ORDER BY m.date, o.match_id, o.market, o.stage, o.bookmaker, o.selection"),
]

NOTES = [
    "Fuentes: LaLiga = football-data.co.uk; Euroliga = API oficial (api-live.euroleague.net); "
    "Liga Endesa = acb.com.",
    "Sólo LaLiga tiene cuotas: ni la Euroliga ni la ACB las publican en sus fuentes oficiales.",
    "result: H = gana local, D = empate, A = gana visitante. is_home: 1 = local, 0 = visitante.",
    "Horas: LaLiga en hora del Reino Unido (así las da la fuente); baloncesto en UTC.",
    "Cuotas: stage 'pre' = recogidas antes del partido; 'closing' = cuota de cierre. "
    "market 1X2 / OU (más-menos 2,5 goles) / AH (hándicap asiático; line = hándicap del local).",
    "Baloncesto: fg2/fg3 = tiros de 2/3 (m anotados, a intentados); ft = tiros libres; "
    "oreb/dreb/treb = rebotes; rating/val = valoración; dunks sólo en Liga Endesa.",
    "status 'not_played': partido cancelado (p. ej. clubes rusos en la Euroliga 2021-22).",
]


def _cell(ws, value, head=False):
    c = WriteOnlyCell(ws, value=value)
    c.font = HEAD_FONT if head else FONT
    if head:
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center")
    return c


def export(db=DEFAULT_DB, out=None) -> Path:
    db = Path(db)
    out = Path(out) if out else db.with_suffix(".xlsx")
    con = sqlite3.connect(db)
    wb = Workbook(write_only=True)
    info = wb.create_sheet("Léeme")
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 14
    info.column_dimensions["C"].width = 110
    info.append([_cell(info, "Base de datos deportiva: LaLiga, Euroliga y Liga Endesa", head=True)])
    info.append([])
    info.append([_cell(info, h, head=True) for h in ("Hoja", "Filas", "Contenido")])
    summary_row = {}
    sheets = []
    for name, desc, sql in SHEETS:
        ws = wb.create_sheet(name)
        cur = con.execute(sql)
        cols = [d[0] for d in cur.description]
        widths = [max(10, len(c) + 2) for c in cols]
        ws.freeze_panes = "A2"
        rows = cur.fetchall()
        for r in rows[:200]:
            for i, v in enumerate(r):
                widths[i] = min(40, max(widths[i], len(str(v)) + 2 if v is not None else 0))
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
        ws.append([_cell(ws, c, head=True) for c in cols])
        for r in rows:
            ws.append([_cell(ws, v) for v in r])
        sheets.append((name, len(rows), desc))
    for name, n, desc in sheets:
        info.append([_cell(info, name), _cell(info, n), _cell(info, desc)])
    info.append([])
    info.append([_cell(info, "Cobertura", head=True), _cell(info, "", head=True),
                 _cell(info, "", head=True)])
    for league, season, n, played, d0, d1 in con.execute(
            "SELECT league, season, COUNT(*), SUM(status='played'), MIN(date), MAX(date) "
            "FROM matches GROUP BY 1, 2 ORDER BY 1, 2"):
        info.append([_cell(info, league), _cell(info, season),
                     _cell(info, f"{n} partidos, {played} jugados ({d0} a {d1})")])
    info.append([])
    info.append([_cell(info, "Notas", head=True), _cell(info, "", head=True),
                 _cell(info, "", head=True)])
    for n in NOTES:
        info.append([_cell(info, ""), _cell(info, ""), _cell(info, n)])
    wb.save(out)
    con.close()
    print(f"Excel escrito en {out}")
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--db", default=str(DEFAULT_DB))
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    export(a.db, a.out)
